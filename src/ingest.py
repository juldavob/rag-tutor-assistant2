"""
ingest.py
---------
Ingesta: extracción de texto desde PDF, DOCX, TXT/MD y HTML.

Cada archivo se convierte en uno o más `Document`:
  - PDF: un Document por página (así se puede citar la página).
  - Resto: un Document por archivo (page = None).

Los encabezados de DOCX y HTML se normalizan a sintaxis Markdown ('#', '##')
para que el chunker pueda cortar por secciones de forma uniforme.
"""

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx", ".html", ".htm"}


@dataclass
class Document:
    source: str            # nombre del archivo (se usa como cita)
    title: str             # título legible del documento
    text: str
    page: Optional[int] = None


def _title_from_text(text: str, fallback: str) -> str:
    match = re.search(r"^#\s+(.+)$", text, flags=re.MULTILINE)
    return match.group(1).strip() if match else fallback


def _read_text_file(path: Path) -> List[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [Document(path.name, _title_from_text(text, path.stem), text)]


def _read_pdf(path: Path) -> List[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(Document(path.name, path.stem, text, page=number))
    return docs


def _read_docx(path: Path) -> List[Document]:
    from docx import Document as DocxDocument

    docx = DocxDocument(str(path))
    lines = []
    for paragraph in docx.paragraphs:
        content = paragraph.text.strip()
        if not content:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        match = re.match(r"heading (\d)", style)
        if match:
            lines.append("#" * int(match.group(1)) + " " + content)
        elif style == "title":
            lines.append("# " + content)
        else:
            lines.append(content)
    text = "\n\n".join(lines)
    return [Document(path.name, _title_from_text(text, path.stem), text)]


def _read_html(path: Path) -> List[Document]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    lines = []
    for element in soup.find_all(["h1", "h2", "h3", "h4", "p", "li"]):
        content = element.get_text(" ", strip=True)
        if not content:
            continue
        if element.name.startswith("h"):
            lines.append("#" * int(element.name[1]) + " " + content)
        else:
            lines.append(content)
    text = "\n\n".join(lines)
    return [Document(path.name, _title_from_text(text, path.stem), text)]


_READERS = {
    ".md": _read_text_file,
    ".txt": _read_text_file,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".html": _read_html,
    ".htm": _read_html,
}


def load_documents(directory: Path) -> List[Document]:
    """Lee todos los archivos soportados de `directory` (orden alfabético)."""
    documents: List[Document] = []
    for path in sorted(Path(directory).rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            documents.extend(_READERS[path.suffix.lower()](path))
    return documents


def corpus_fingerprint(directory: Path) -> str:
    """Hash del contenido del corpus: si cambia un archivo, cambia el índice."""
    digest = hashlib.sha256()
    for path in sorted(Path(directory).rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            digest.update(path.name.encode("utf-8"))
            digest.update(path.read_bytes())
    return digest.hexdigest()
